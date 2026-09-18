import io
import re
import tempfile
from pathlib import Path

import ezdxf
import pandas as pd
import streamlit as st

st.set_page_config(page_title='Leitor de Projeto BR7', layout='wide')
st.title('Leitor de Projeto BR7')
st.caption('Versão web de teste — baseada na lógica da V23')


def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def text_of(e):
    try:
        typ = e.dxftype()
        if typ == 'TEXT':
            return e.dxf.text or ''
        if typ == 'MTEXT':
            return e.text or ''
        if typ in ('ATTRIB', 'ATTDEF'):
            return e.dxf.text or ''
    except Exception:
        pass
    return ''


def xy_of(e):
    try:
        p = e.dxf.insert
        return float(p.x), float(p.y)
    except Exception:
        try:
            p = e.dxf.location
            return float(p.x), float(p.y)
        except Exception:
            return 0.0, 0.0


def read_dxf(data):
    with tempfile.NamedTemporaryFile(suffix='.dxf', delete=False) as f:
        f.write(data)
        path = f.name
    try:
        doc = ezdxf.readfile(path)
    finally:
        Path(path).unlink(missing_ok=True)
    rows = []
    for e in doc.modelspace():
        t = norm(text_of(e))
        if t:
            x, y = xy_of(e)
            rows.append({'Texto': t, 'X': x, 'Y': y})
    return doc, rows


def headers(rows):
    out = []
    pats = [re.compile(r'\bPP\s*[-–—:]\s*([A-Z])\b', re.I),
            re.compile(r'\bMODELO\s*[-–—:]?\s*([A-Z])\b', re.I)]
    for r in rows:
        for p in pats:
            m = p.search(r['Texto'])
            if m:
                out.append({'Modelo': m.group(1).upper(), 'X': r['X'], 'Y': r['Y']})
                break
    if out:
        return out
    xs = sorted(r['X'] for r in rows if 'módulo' in r['Texto'].lower() and ('inicial' in r['Texto'].lower() or 'adicional' in r['Texto'].lower()))
    groups = []
    for x in xs:
        if not groups or abs(x - groups[-1][-1]) > 1000:
            groups.append([x])
        else:
            groups[-1].append(x)
    centers = [sum(g) / len(g) for g in groups]
    return [{'Modelo': chr(ord('A') + i), 'X': x, 'Y': 0.0} for i, x in enumerate(centers)]


def nearest(r, hs):
    if not hs:
        return 'A'
    return min(hs, key=lambda h: (abs(r['X']-h['X']), abs(r['Y']-h['Y'])))['Modelo']


def parse_models(rows, hs):
    cfgs = {}
    for r in rows:
        t = norm(r['Texto']); low = t.lower()
        if 'módulo' not in low or not ('inicial' in low or 'adicional' in low):
            continue
        m = re.search(r'(\d+)\s*m[oó]dulo(?:s)?\b', t, re.I)
        if not m:
            continue
        mod = nearest(r, hs)
        cfg = cfgs.setdefault(mod, {'initial': 0, 'additional': 0, 'total': 0})
        cfg['initial' if 'inicial' in low else 'additional'] += int(m.group(1))
    for cfg in cfgs.values():
        cfg['total'] = cfg['initial'] + cfg['additional']
    return cfgs


def parse_codes(rows, hs):
    pat = re.compile(r'\b\d{3}(?:\.\d{2}\.\d\.\d{4}|\s+\d{5})\b')
    mods, codes = [], []
    for r in rows:
        t = r['Texto']; low = t.lower()
        if 'módulo' in low and 'inicial' in low:
            mods.append({'modelo': nearest(r, hs), 'tipo': 'Módulo simples inicial', 'X': r['X'], 'Y': r['Y']})
        elif 'módulo' in low and 'adicional' in low:
            mods.append({'modelo': nearest(r, hs), 'tipo': 'Módulo simples adicional', 'X': r['X'], 'Y': r['Y']})
        m = pat.search(t)
        if m:
            codes.append({'code': m.group(0), 'X': r['X'], 'Y': r['Y']})
    result = {}
    for mod in mods:
        available = [c for c in codes if c['code'] not in result.values()]
        if available:
            best = min(available, key=lambda c: abs(c['X']-mod['X']) + 2*abs(c['Y']-mod['Y']))
            result[(mod['modelo'], mod['tipo'])] = best['code']
    return result


def parse_plans(rows, hs, cfgs):
    data = {m: {'tipo': 'Plano palete', 'material': 'Estrutura padrão para palete', 'planos': 0, 'reforcos': 0} for m in cfgs}
    for r in rows:
        t = r['Texto']; low = t.lower(); mod = nearest(r, hs)
        if mod not in data:
            continue
        if any(k in low for k in ('picking', 'bandeja', 'bandejas', 'mdf', 'mdp', 'madeira')):
            if any(k in low for k in ('mdf', 'mdp', 'madeira')):
                data[mod]['tipo'] = 'Plano picking – madeira (MDF)'; data[mod]['material'] = 'MDF / madeira'
            else:
                data[mod]['tipo'] = 'Plano picking – aço (bandejas)'; data[mod]['material'] = 'Bandeja em aço galvanizado'
        mp = re.search(r'\b0*(\d+)\s*(?:planos?|pares?)\b', t, re.I)
        if mp:
            data[mod]['planos'] = int(mp.group(1))
        mr = re.search(r'\b0*(\d+)\s*refor[cç]o(?:s)?\b', t, re.I)
        if mr:
            data[mod]['reforcos'] = int(mr.group(1))
    return data


def parse_materials(rows, hs, cfgs):
    data = {m: {} for m in cfgs}
    for r in rows:
        t = r['Texto']; low = t.lower(); mod = nearest(r, hs)
        if mod not in data:
            continue
        if low.startswith('coluna:'):
            data[mod]['Coluna'] = t.split(':', 1)[1].strip()
        elif low.startswith('travessa e diagonal:'):
            v = t.split(':', 1)[1].strip()
            data[mod].update({'Travessa': v, 'Diagonal maior': v, 'Diagonal menor': v})
        elif low.startswith('travessa:'):
            data[mod]['Travessa'] = t.split(':', 1)[1].strip()
        elif low.startswith('diagonal maior:'):
            data[mod]['Diagonal maior'] = t.split(':', 1)[1].strip()
        elif low.startswith('diagonal menor:'):
            data[mod]['Diagonal menor'] = t.split(':', 1)[1].strip()
        elif low.startswith('longarina:'):
            data[mod]['Longarina'] = t.split(':', 1)[1].strip()
    return data


def parse_dimensions(rows, hs, cfgs):
    out = {m: None for m in cfgs}
    p = re.compile(r'(\d{1,2}(?:[.\s]\d{3})?)\s*[x×]\s*(\d{1,2}(?:[.\s]\d{3})?)\s*[x×]\s*(\d{1,2}(?:[.\s]\d{3})?)\s*mm', re.I)
    for r in rows:
        m = p.search(r['Texto'])
        if m:
            out[nearest(r, hs)] = [int(re.sub(r'[^0-9]', '', x)) for x in m.groups()]
    return out


def build(rows):
    hs = headers(rows); cfgs = parse_models(rows, hs)
    if not cfgs:
        cfgs = {'A': {'initial': 0, 'additional': 0, 'total': 0}}
        hs = [{'Modelo': 'A', 'X': 0.0, 'Y': 0.0}]
    codes = parse_codes(rows, hs); plans = parse_plans(rows, hs, cfgs); mats = parse_materials(rows, hs, cfgs); dims = parse_dimensions(rows, hs, cfgs)
    conjuntos, planos, comp = [], [], []
    for modelo, cfg in cfgs.items():
        h = dims.get(modelo)[2] if dims.get(modelo) else 0
        p = plans[modelo]; mat = mats.get(modelo, {})
        planos.append({'Modelo': modelo, 'Tipo de plano': p['tipo'], 'Material': p['material'], 'Planos por conjunto': p['planos'], 'Reforços por plano': p['reforcos'] or 0})
        for tipo, qtd, mp in [('Módulo simples inicial', cfg['initial'], 2), ('Módulo simples adicional', cfg['additional'], 1)]:
            conjuntos.append({'Modelo': modelo, 'Código': codes.get((modelo, tipo), 'NÃO IDENTIFICADO'), 'Conjunto': tipo, 'Quantidade': qtd, 'Altura de projeto (mm)': h,
                              'Coluna': mat.get('Coluna', 'NÃO IDENTIFICADO'), 'Travessa': mat.get('Travessa', 'NÃO IDENTIFICADO'),
                              'Diagonal maior': mat.get('Diagonal maior', 'NÃO IDENTIFICADO'), 'Diagonal menor': mat.get('Diagonal menor', 'NÃO IDENTIFICADO'), 'Longarina': mat.get('Longarina', 'NÃO IDENTIFICADO')})
            qm = qtd * mp; lp = p['planos'] * 2
            comp.append({'Modelo': modelo, 'Código': codes.get((modelo, tipo), 'NÃO IDENTIFICADO'), 'Conjunto': tipo, 'Quantidade de conjuntos': qtd,
                         'Montantes por conjunto': mp, 'Montantes totais': qm, 'Altura (mm)': h, 'Travessas totais': qm*2,
                         'Diagonal menor total': qm*3, 'Diagonal maior total': qm*7, 'Longarinas por conjunto': lp, 'Longarinas totais': qtd*lp,
                         'Coluna - material/cor': mat.get('Coluna', 'NÃO IDENTIFICADO'), 'Travessa - material/cor': mat.get('Travessa', 'NÃO IDENTIFICADO'),
                         'Diagonal maior - material/cor': mat.get('Diagonal maior', 'NÃO IDENTIFICADO'), 'Diagonal menor - material/cor': mat.get('Diagonal menor', 'NÃO IDENTIFICADO'),
                         'Longarina - material/cor': mat.get('Longarina', 'NÃO IDENTIFICADO'), 'Observação': 'Total da linha; não é quantidade por peça.'})
    return pd.DataFrame(conjuntos), pd.DataFrame(planos), pd.DataFrame(comp)


def xlsx(df):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Alignment
    out = io.BytesIO(); wb = Workbook(); ws = wb.active; ws.title = 'Tabela'
    ws.append([str(c) for c in df.columns])
    for row in df.itertuples(index=False, name=None): ws.append(list(row))
    if ws.max_row >= 2:
        ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
        t = Table(displayName='TabelaExportada', ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(t)
    for col in range(1, ws.max_column+1):
        letter = get_column_letter(col); head = str(ws.cell(1, col).value or '')
        if head.lower() == 'observação':
            ws.column_dimensions[letter].width = 60
        else:
            w = max(len(str(ws.cell(r, col).value or '')) for r in range(1, ws.max_row+1))
            ws.column_dimensions[letter].width = min(max(w+2, 12), 32)
    wb.save(out); out.seek(0); return out.getvalue()

uploaded = st.file_uploader('Envie um arquivo DXF', type=['dxf'])
if uploaded:
    try:
        _, rows = read_dxf(uploaded.getvalue())
        conjuntos_df, planos_df, comp_df = build(rows)
        st.success('Arquivo lido com sucesso.')

        st.subheader('Conjuntos identificados')
        st.dataframe(conjuntos_df, use_container_width=True, hide_index=True)

        st.subheader('Planos')
        st.dataframe(planos_df, use_container_width=True, hide_index=True)

        st.subheader('Composição técnica por conjunto')
        st.dataframe(comp_df, use_container_width=True, hide_index=True)
        st.download_button(
            'Baixar composição por conjunto Excel (.xlsx)',
            xlsx(comp_df),
            'composicao_por_conjunto_br7.xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
    except Exception as exc:
        st.error('Não foi possível ler o DXF.')
        st.exception(exc)
